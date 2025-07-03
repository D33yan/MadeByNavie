from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# Title
doc.add_heading('DIVINE NNAJI', 0)
doc.add_paragraph('Front-end Developer')
doc.add_paragraph('Bwari, Abuja, Nigeria')
doc.add_paragraph('Phone: +234 810 689 0380')
doc.add_paragraph('Email: dnnaji26@gmail.com')
doc.add_paragraph('LinkedIn: https://www.linkedin.com/in/divine-nnaji-858a53283')
doc.add_paragraph('GitHub: https://github.com/D33yan')

doc.add_heading('Professional Summary', level=1)
doc.add_paragraph(
    'Results-driven Front-end Developer with proven experience building responsive, cross-platform web and mobile applications. '
    'Adept at modern JavaScript frameworks, UI/UX design, and collaborating in agile teams. Passionate about delivering high-quality, user-centric solutions and eager to contribute to innovative projects.'
)

doc.add_heading('Technical Skills', level=1)
doc.add_paragraph('Languages & Frameworks: React, React Native, Next.js, JavaScript (ES6+), PHP, Python')
doc.add_paragraph('Styling & UI: Tailwind CSS, Material Design, Shadcn UI')
doc.add_paragraph('Databases & Backend: Firebase, Supabase, MongoDB, RESTful APIs')
doc.add_paragraph('Tools & Platforms: Git, GitHub, Android Studio, VS Code')

doc.add_heading('Soft Skills', level=1)
doc.add_paragraph('• Problem Solving\n• Team Collaboration\n• Strong Communication\n• Attention to Detail\n• Time Management\n• Adaptability')

doc.add_heading('Projects', level=1)
projects = [
    {
        'name': 'Auction App',
        'type': 'Personal Project',
        'desc': 'Developed a mobile auction app using React Native and JavaScript. Designed an intuitive interface for managing auction plans and tasks, with persistent async data storage.',
        'github': 'https://github.com/D33yan/rebid-app'
    },
    {
        'name': 'Typhoid Checker',
        'type': 'Personal Project',
        'desc': 'Built a health diagnostic app to help users assess typhoid symptoms using React Native. Integrated user-friendly forms and provided actionable health recommendations.',
        'github': 'https://github.com/D33yan/typhoid-checker'
    },
    {
        'name': 'Fitness Tracker',
        'type': 'Personal Project',
        'desc': 'Created a cross-platform fitness tracking app with React Native. Enabled users to log workouts, track progress, and visualize fitness data with engaging UI components.',
        'github': 'https://github.com/D33yan/fitnesstracker'
    },
    {
        'name': 'E-commerce Website',
        'type': 'Personal Project',
        'desc': 'Built a shopping and blog-integrated website with React and Supabase. Implemented product management and cart functionality using React Context.',
        'github': 'https://github.com/D33yan/afabric-ecommercestore'
    },
    {
        'name': 'Academic Journal Website',
        'type': 'Paid Project',
        'desc': 'Designed and developed an academic journal and trading portfolio site using Next.js, Tailwind CSS, and Shadcn UI. Built interfaces for displaying research articles, trades, and services, with integrated contact features.',
        'github': 'https://github.com/D33yan/firebase-blog'
    }
]
for p in projects:
    para = doc.add_paragraph()
    run = para.add_run(f"{p['name']} ({p['type']})\n")
    run.bold = True
    para.add_run(f"- {p['desc']}\n")
    para.add_run(f"- GitHub: {p['github']}")

doc.add_heading('Certifications', level=1)
doc.add_paragraph(
    'Meta React Developer Certification – Meta, 2023\n'
    'UI/UX Design Professional – Google, 2023\n'
    'Early Code App Development Certification – Early Code Academy, 2022\n'
    'Google Associate Android Developer – Google, 2022\n'
    'AWS Certified Cloud Practitioner – Amazon Web Services, 2022\n'
    'Microsoft Certified: Azure Fundamentals – Microsoft, 2022\n'
    'Coursera Full-Stack Web Development – Coursera / HKUST, 2021\n'
    'Complete Next.js & JavaScript Certificate of Completion – Udemy, 2021'
)

doc.add_heading('Languages', level=1)
doc.add_paragraph('English (Fluent)')

doc.save('Divine_Nnaji_CV.docx') 